import openpyxl

path = r'Z:\6. PREGÃO ELETRONICO e DISP. ELETRÔNICA.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print("Sheets:", wb.sheetnames)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        # Check if 38 is in any cell
        for cell in row:
            if cell is not None:
                val = str(cell).strip()
                if val == '38' or '38/2026' in val or '00038/2026' in val or '38/2025' in val or '38/2024' in val:
                    print(f"[{sheet}] Line {idx}: {' | '.join(str(c) for c in row if c is not None)}")
                    break
