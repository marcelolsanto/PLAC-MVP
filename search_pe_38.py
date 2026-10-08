import openpyxl

path = r'Z:\6. PREGÃO ELETRONICO e DISP. ELETRÔNICA.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['CONTROLE PEs e DEs']

header = [cell for cell in next(ws.iter_rows(values_only=True))]
print("Colunas:", header[:15])

for idx, row in enumerate(ws.iter_rows(values_only=True)):
    pe_num = str(row[8]) if len(row) > 8 else ''
    processo = str(row[7]) if len(row) > 7 else ''
    objeto = str(row[9]) if len(row) > 9 else ''
    links = [str(c) for c in row if c and 'cnetmobile' in str(c)]
    
    if '38' in pe_num or '38' in processo:
        print(f"Row {idx} -> PE: {pe_num} | Proc: {processo} | Obj: {objeto[:60]} | Link: {links}")
