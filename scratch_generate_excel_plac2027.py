import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

print("Gerando Relatorio_Esteira_PLAC_2027_Vigentes_GCC_Demandantes.xlsx...")

wb_out = openpyxl.Workbook()
wb_out.remove(wb_out.active)

# Estilos
navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
blue_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
teal_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
alert_red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
alert_yellow_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
alert_green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

font_white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
font_regular = Font(name="Calibri", size=10, color="1E293B")
font_red = Font(name="Calibri", size=10, bold=True, color="991B1B")
font_green = Font(name="Calibri", size=10, bold=True, color="166534")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

# 1. Carregar Kanban
wb_k = openpyxl.load_workbook('Z:/PLAC-MVP/BASE_PROCESSOS_PLANEJAMENTO_KANBAN_SIGA_2026_ATUALIZADA.xlsx', data_only=True)
ws_k = wb_k['INVENTARIO_GERAL']
processos = []
for r in list(ws_k.iter_rows(values_only=True))[1:]:
    if r[1]:
        processos.append({
            'proc': str(r[1]).strip(),
            'cod_plac': str(r[2]).strip() if r[2] else '',
            'origem': str(r[3]).strip() if r[3] else '',
            'dir': str(r[4]).strip() if r[4] else '',
            'ger': str(r[5]).strip() if r[5] else '',
            'objeto': str(r[6]).strip() if r[6] else '',
            'val': float(r[7]) if r[7] and isinstance(r[7], (int, float)) else 0.0,
            'fase': str(r[9]).strip() if r[9] else '',
            'setor': str(r[10]).strip() if r[10] else '',
            'custodiante': str(r[11]).strip() if r[11] else '',
            'dias': int(r[12]) if r[12] and isinstance(r[12], (int, float)) else 0,
            'sla': int(r[13]) if r[13] and isinstance(r[13], (int, float)) else 0,
            'desvio': int(r[14]) if r[14] and isinstance(r[14], (int, float)) else 0,
            'status': str(r[15]).strip() if r[15] else '',
            'acao': str(r[17]).strip() if len(r) > 17 and r[17] else '',
            'prox_doc': str(r[19]).strip() if len(r) > 19 and r[19] else ''
        })

# 2. Carregar Contratos Vigentes
wb_v = openpyxl.load_workbook('Z:/PLAC-MVP/AUDITORIA_CONTRATOS_VIGENTES_2026_SIGA.xlsx', data_only=True)
ws_v = wb_v['CONTRATOS_VIGENTES_2026']
contratos = []
for r in list(ws_v.iter_rows(values_only=True))[1:]:
    if r[1]:
        contratos.append({
            'ctr': str(r[1]).strip(),
            'proc': str(r[2]).strip() if r[2] else '',
            'dir': str(r[3]).strip() if r[3] else '',
            'area': str(r[4]).strip() if r[4] else '',
            'modalidade': str(r[5]).strip() if r[5] else '',
            'ini': str(r[6])[:10] if r[6] else '',
            'fim': str(r[7])[:10] if r[7] else '',
            'val': float(r[8]) if r[8] and isinstance(r[8], (int, float)) else 0.0,
            'cod_plac': str(r[10]).strip() if r[10] else '',
            'forn': str(r[11]).strip() if r[11] else '',
            'objeto': str(r[12]).strip() if r[12] else ''
        })

# ABA 1: PROCESSOS NA GCC (SETOR 2600)
ws1 = wb_out.create_sheet(title="1_Planejamento_Mesa_GCC")
ws1.views.sheetView[0].showGridLines = True
headers1 = ["Item", "Processo SIGA", "Código PLAC", "Diretoria", "Gerência Demandante", "Objeto", "Valor Estimado (R$)", "Fase Kanban", "Setor Atual", "Custodiante", "Dias no Setor", "SLA (dias)", "Desvio", "Status", "Ação Recomendada", "Próximo Documento"]

ws1.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=c_idx, value=h)
    cell.font = font_white_bold
    cell.fill = blue_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

procs_gcc = [p for p in processos if '2600' in p['setor'] or 'GCC' in p['setor'].upper()]
for idx, p in enumerate(procs_gcc, 2):
    ws1.row_dimensions[idx].height = 19
    ws1.cell(row=idx, column=1, value=idx-1).alignment = Alignment(horizontal="center")
    ws1.cell(row=idx, column=2, value=p['proc'])
    ws1.cell(row=idx, column=3, value=p['cod_plac'])
    ws1.cell(row=idx, column=4, value=p['dir'])
    ws1.cell(row=idx, column=5, value=p['ger'])
    ws1.cell(row=idx, column=6, value=p['objeto'])
    c_val = ws1.cell(row=idx, column=7, value=p['val'])
    c_val.number_format = 'R$ #,##0.00'
    ws1.cell(row=idx, column=8, value=p['fase'])
    ws1.cell(row=idx, column=9, value=p['setor'])
    ws1.cell(row=idx, column=10, value=p['custodiante'])
    ws1.cell(row=idx, column=11, value=p['dias']).alignment = Alignment(horizontal="center")
    ws1.cell(row=idx, column=12, value=p['sla']).alignment = Alignment(horizontal="center")
    ws1.cell(row=idx, column=13, value=p['desvio']).alignment = Alignment(horizontal="center")
    c_st = ws1.cell(row=idx, column=14, value=p['status'])
    c_st.alignment = Alignment(horizontal="center")
    if "ATENÇÃO" in p['status'].upper():
        c_st.fill = alert_yellow_fill
    else:
        c_st.fill = alert_green_fill
        c_st.font = font_green
    ws1.cell(row=idx, column=15, value=p['acao'])
    ws1.cell(row=idx, column=16, value=p['prox_doc'])
    for c in range(1, 17):
        ws1.cell(row=idx, column=c).border = thin_border

# ABA 2: PROCESSOS NAS ÁREAS DEMANDANTES
ws2 = wb_out.create_sheet(title="2_Planejamento_Demandantes")
ws2.views.sheetView[0].showGridLines = True
headers2 = ["Item", "Processo SIGA", "Código PLAC", "Diretoria", "Gerência Demandante", "Objeto", "Valor Estimado (R$)", "Fase Kanban", "Setor Atual", "Custodiante", "Dias no Setor", "SLA (dias)", "Desvio", "Status", "Ação em Andamento", "Próximo Documento"]

ws2.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=1, column=c_idx, value=h)
    cell.font = font_white_bold
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

procs_dem = [p for p in processos if p not in procs_gcc and not ('CONJUR' in p['setor'].upper() or '1600' in p['setor'])]
for idx, p in enumerate(procs_dem, 2):
    ws2.row_dimensions[idx].height = 19
    ws2.cell(row=idx, column=1, value=idx-1).alignment = Alignment(horizontal="center")
    ws2.cell(row=idx, column=2, value=p['proc'])
    ws2.cell(row=idx, column=3, value=p['cod_plac'])
    ws2.cell(row=idx, column=4, value=p['dir'])
    ws2.cell(row=idx, column=5, value=p['ger'])
    ws2.cell(row=idx, column=6, value=p['objeto'])
    c_val = ws2.cell(row=idx, column=7, value=p['val'])
    c_val.number_format = 'R$ #,##0.00'
    ws2.cell(row=idx, column=8, value=p['fase'])
    ws2.cell(row=idx, column=9, value=p['setor'])
    ws2.cell(row=idx, column=10, value=p['custodiante'])
    ws2.cell(row=idx, column=11, value=p['dias']).alignment = Alignment(horizontal="center")
    ws2.cell(row=idx, column=12, value=p['sla']).alignment = Alignment(horizontal="center")
    ws2.cell(row=idx, column=13, value=p['desvio']).alignment = Alignment(horizontal="center")
    c_st = ws2.cell(row=idx, column=14, value=p['status'])
    c_st.alignment = Alignment(horizontal="center")
    if "CRÍTICO" in p['status'].upper() or "GARGALO" in p['status'].upper():
        c_st.fill = alert_red_fill
        c_st.font = font_red
    elif "ATENÇÃO" in p['status'].upper():
        c_st.fill = alert_yellow_fill
    else:
        c_st.fill = alert_green_fill
        c_st.font = font_green
    ws2.cell(row=idx, column=15, value=p['acao'])
    ws2.cell(row=idx, column=16, value=p['prox_doc'])
    for c in range(1, 17):
        ws2.cell(row=idx, column=c).border = thin_border

# ABA 3: CONTRATOS VIGENTES & TRANSIÇÃO PLAC 2027
ws3 = wb_out.create_sheet(title="3_Vigentes_Gestao_Contratual")
ws3.views.sheetView[0].showGridLines = True
headers3 = ["Item", "Nº Contrato", "Processo SIGA", "Diretoria", "Área Demandante", "Fornecedor", "Objeto Contratado", "Valor Contrato (R$)", "Início Vigência", "Fim Vigência", "Ano Vencimento", "Impacto no PLAC 2027", "Ação de Gestão Contratual Requerida"]

ws3.row_dimensions[1].height = 26
for c_idx, h in enumerate(headers3, 1):
    cell = ws3.cell(row=1, column=c_idx, value=h)
    cell.font = font_white_bold
    cell.fill = teal_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for idx, c in enumerate(contratos, 2):
    ws3.row_dimensions[idx].height = 19
    ws3.cell(row=idx, column=1, value=idx-1).alignment = Alignment(horizontal="center")
    ws3.cell(row=idx, column=2, value=c['ctr'])
    ws3.cell(row=idx, column=3, value=c['proc'])
    ws3.cell(row=idx, column=4, value=c['dir'])
    ws3.cell(row=idx, column=5, value=c['area'])
    ws3.cell(row=idx, column=6, value=c['forn'])
    ws3.cell(row=idx, column=7, value=c['objeto'])
    c_val = ws3.cell(row=idx, column=8, value=c['val'])
    c_val.number_format = 'R$ #,##0.00'
    ws3.cell(row=idx, column=9, value=c['ini']).alignment = Alignment(horizontal="center")
    ws3.cell(row=idx, column=10, value=c['fim']).alignment = Alignment(horizontal="center")
    
    ano_fim = c['fim'][:4] if len(c['fim']) >= 4 else 'N/D'
    ws3.cell(row=idx, column=11, value=ano_fim).alignment = Alignment(horizontal="center")
    
    c_imp = ws3.cell(row=idx, column=12)
    c_acao = ws3.cell(row=idx, column=13)
    c_imp.alignment = Alignment(horizontal="center")
    
    if ano_fim == '2026':
        c_imp.value = "Crítico: Vence em 2026 (Transição Imediata)"
        c_imp.fill = alert_red_fill
        c_imp.font = font_red
        c_acao.value = "Disparar T-120 para Aditivo ou Nova Licitação no Q1/2027"
    elif ano_fim == '2027':
        c_imp.value = "Mandatório no PLAC 2027 (Renovação / F2=5)"
        c_imp.fill = alert_yellow_fill
        c_acao.value = "Cadastrar no Levantamento 2027 com antecedência do Cronograma Reverso"
    else:
        c_imp.value = "Plurianual (Monitoramento Regular)"
        c_imp.fill = alert_green_fill
        c_imp.font = font_green
        c_acao.value = "Rotina de fiscalização mensal e apostilamento anual na GCC"

    for col in range(1, 14):
        ws3.cell(row=idx, column=col).border = thin_border

# Autoajuste de colunas
for ws in [ws1, ws2, ws3]:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 1 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 50)

out_path = "Z:/PLAC-MVP/Relatorio_Esteira_PLAC_2027_Vigentes_GCC_Demandantes.xlsx"
wb_out.save(out_path)
print(f"Planilha salva com sucesso em: {out_path}!")
