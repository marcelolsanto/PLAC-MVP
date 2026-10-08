import openpyxl

path = r'Z:\6. PREGÃO ELETRONICO e DISP. ELETRÔNICA.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        txt = " | ".join(str(c) for c in row if c is not None)
        if '38/2026' in txt or '38' in txt and '2026' in txt:
            print(f"[{sheet}] Line {idx}: {txt[:200]}")
